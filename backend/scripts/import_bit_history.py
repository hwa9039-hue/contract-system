"""BIT 이력관리 — 기존 엑셀 기록을 bit_history_rows 로 일괄 입력하는 일회성 스크립트.

사용법 (backend 폴더에서):
  # 1) 미리보기 (DB 변경 없음)
  .\\.venv\\Scripts\\python.exe scripts\\import_bit_history.py
  # 2) 실제 입력
  .\\.venv\\Scripts\\python.exe scripts\\import_bit_history.py --commit

옵션:
  --file <경로>    대상 엑셀 (기본: ../data/input 에서 가장 최근 .xlsx)
  --sheet <이름>   대상 시트 (기본: 첫 번째 시트)
  --commit         실제 insert 실행 (미지정 시 미리보기만)

중복 검사 없이 엑셀에 있는 모든 행을 추가한다.
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from uuid import uuid4

BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from openpyxl import load_workbook  # noqa: E402

from app.database import get_connection  # noqa: E402

DEFAULT_INPUT_DIR = BACKEND_DIR.parent / "data" / "input"

# DB 컬럼 → 엑셀 헤더 후보. 정규화(공백·괄호 제거) 후 비교한다.
COLUMN_HEADER_CANDIDATES = {
    "seq_no": ["순번", "번호", "no"],
    "client": ["발주처", "발주기관", "고객사", "수요기관"],
    "department": ["담당부서", "부서", "담당부서명"],
    "contract_method": ["계약방식", "계약방법"],
    "contract_class": ["계약분류", "물품분류번호"],
    "ident_no": ["식별번호", "식별no", "관리번호"],
    "contract_date": ["계약일자", "계약일"],
    "due_date": ["납기일", "납기일자", "납기", "납품기한"],
    "project_name": ["사업명", "사업명칭", "과업명", "건명"],
    "contract_amount": ["계약금액", "금액", "계약액"],
    "quantity": ["수량"],
    "board_applied": ["보드적용", "통신모뎀", "모뎀"],
    "program_item": ["프로그램"],
    "manufacturing": ["제작"],
    "shipping_inspection": ["출하검사", "출하"],
    "note1": ["비고1"],
    "module_item": ["모듈"],
    "module_array": ["배열"],
    "module_kind": ["종류"],
    "etc_item": ["기타"],
    "project_complete": ["사업완료", "완료"],
    "defect": ["불량발생", "불량"],
    "note2": ["비고2"],
}

DATE_COLUMNS = {"contract_date", "due_date"}
AMOUNT_COLUMNS = {"contract_amount"}
# 수식 오류(#VALUE! 등)를 같은 열의 최빈값으로 보정할 컬럼
ERROR_FIXUP_COLUMNS = {"contract_class"}
ALL_COLUMNS = list(COLUMN_HEADER_CANDIDATES.keys())


def normalize_header(value) -> str:
    """'비고 (1)' → '비고1', '계약 일자' → '계약일자'"""
    text = "" if value is None else str(value)
    return re.sub(r"[\s()\[\]{}·.\-_/:]+", "", text).lower()


def build_header_lookup() -> dict[str, str]:
    lookup: dict[str, str] = {}
    for column, candidates in COLUMN_HEADER_CANDIDATES.items():
        for candidate in candidates:
            lookup[normalize_header(candidate)] = column
    return lookup


HEADER_LOOKUP = build_header_lookup()


def map_header_row(cells: list) -> dict[int, str]:
    """헤더 행 → {열 인덱스: DB 컬럼}. 번호 없는 '비고'는 나온 순서대로 note1, note2."""
    mapping: dict[int, str] = {}
    plain_note_seen = 0
    for index, cell in enumerate(cells):
        key = normalize_header(cell)
        if not key:
            continue
        if key == "비고":
            plain_note_seen += 1
            column = "note1" if plain_note_seen == 1 else "note2"
            if column not in mapping.values():
                mapping[index] = column
            continue
        column = HEADER_LOOKUP.get(key)
        if column and column not in mapping.values():
            mapping[index] = column
    return mapping


def find_header_row(worksheet, max_scan: int = 30) -> tuple[int, dict[int, str]]:
    """상단 병합·제목 줄을 건너뛰고, 우리 컬럼이 가장 많이 잡히는 행을 헤더로 본다."""
    best: tuple[int, dict[int, str]] = (-1, {})
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=max_scan, values_only=True)
    ):
        mapping = map_header_row(list(row))
        if len(mapping) > len(best[1]):
            best = (row_index + 1, mapping)
    return best


def cell_to_date_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return ""
    matched = re.match(r"^(\d{4})\s*[-./년]\s*(\d{1,2})\s*[-./월]\s*(\d{1,2})", text)
    if matched:
        year, month, day = (int(part) for part in matched.groups())
        try:
            return date(year, month, day).isoformat()
        except ValueError:
            return text
    return text


def cell_to_amount_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return f"{int(round(value)):,}"
    digits = re.sub(r"[^\d]", "", str(value))
    if not digits:
        return str(value).strip()
    return f"{int(digits):,}"


def cell_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def is_excel_error_text(value: str) -> bool:
    return value.startswith("#") and value.endswith("!") or value in {"#N/A", "#VALUE!", "#REF!"}


def fix_error_cells(rows: list[dict]) -> int:
    """계약분류처럼 전 행이 같은 값이어야 하는 열의 수식 오류를 최빈값으로 보정한다."""
    fixed = 0
    for column in ERROR_FIXUP_COLUMNS:
        counter = Counter(
            row[column] for row in rows if row[column] and not is_excel_error_text(row[column])
        )
        if not counter:
            continue
        most_common = counter.most_common(1)[0][0]
        for row in rows:
            if is_excel_error_text(row[column]) or not row[column]:
                row[column] = most_common
                fixed += 1
    return fixed


def parse_rows(worksheet, mapping: dict[int, str], header_row: int) -> list[dict]:
    rows: list[dict] = []
    for raw in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        values = {column: "" for column in ALL_COLUMNS}
        for index, column in mapping.items():
            if index >= len(raw):
                continue
            cell = raw[index]
            if column in DATE_COLUMNS:
                values[column] = cell_to_date_text(cell)
            elif column in AMOUNT_COLUMNS:
                values[column] = cell_to_amount_text(cell)
            else:
                values[column] = cell_to_text(cell)

        if not any(values[column] for column in ALL_COLUMNS):
            continue
        # 소계·합계 같은 요약 줄은 건너뛴다
        if not values["project_name"] and not values["client"]:
            continue
        rows.append(values)
    return rows


def resolve_target_file(arg_path: str | None) -> Path:
    if arg_path:
        path = Path(arg_path).expanduser()
        if not path.is_absolute():
            path = (Path.cwd() / path).resolve()
        if not path.exists():
            raise SystemExit(f"파일을 찾을 수 없습니다: {path}")
        return path

    if not DEFAULT_INPUT_DIR.exists():
        raise SystemExit(f"기본 폴더가 없습니다: {DEFAULT_INPUT_DIR}\n--file 로 경로를 지정하세요.")
    candidates = sorted(
        (p for p in DEFAULT_INPUT_DIR.glob("*.xls*") if not p.name.startswith("~$")),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise SystemExit(
            f"{DEFAULT_INPUT_DIR} 에 엑셀 파일이 없습니다. 파일을 넣거나 --file 로 지정하세요."
        )
    return candidates[0]


INSERT_SQL = """
insert into bit_history_rows (
  id, sort_order, seq_no,
  client, department, contract_method, contract_class, ident_no, contract_date, due_date,
  project_name, contract_amount, quantity,
  board_applied, program_item, manufacturing, shipping_inspection, note1,
  module_item, module_array, module_kind, etc_item, project_complete, defect, note2,
  created_at, updated_at
)
values (
  %(id)s, %(sort_order)s, %(seq_no)s,
  %(client)s, %(department)s, %(contract_method)s, %(contract_class)s, %(ident_no)s, %(contract_date)s, %(due_date)s,
  %(project_name)s, %(contract_amount)s, %(quantity)s,
  %(board_applied)s, %(program_item)s, %(manufacturing)s, %(shipping_inspection)s, %(note1)s,
  %(module_item)s, %(module_array)s, %(module_kind)s, %(etc_item)s, %(project_complete)s, %(defect)s, %(note2)s,
  now(), now()
)
"""


def insert_rows(rows: list[dict]) -> int:
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute("select coalesce(max(sort_order), 0) as m from bit_history_rows")
            fetched = cursor.fetchone() or {}
            base = int(fetched.get("m") or 0)
            for offset, values in enumerate(rows, start=1):
                cursor.execute(
                    INSERT_SQL,
                    {**values, "id": str(uuid4()), "sort_order": base + offset},
                )
        connection.commit()
    return len(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="BIT 이력관리 엑셀 일괄 입력")
    parser.add_argument("--file", dest="file", default=None)
    parser.add_argument("--sheet", dest="sheet", default=None)
    parser.add_argument("--commit", dest="commit", action="store_true")
    args = parser.parse_args()

    path = resolve_target_file(args.file)
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet_name = args.sheet or workbook.sheetnames[0]
    if sheet_name not in workbook.sheetnames:
        raise SystemExit(f"시트를 찾을 수 없습니다: {sheet_name} (있는 시트: {workbook.sheetnames})")
    worksheet = workbook[sheet_name]

    header_row, mapping = find_header_row(worksheet)
    if len(mapping) < 5:
        raise SystemExit(
            "헤더를 인식하지 못했습니다. 엑셀 헤더가 '발주처 / 담당부서 / … / 비고' 형태인지 확인하세요.\n"
            f"인식된 컬럼: {sorted(mapping.values())}"
        )

    rows = parse_rows(worksheet, mapping, header_row)
    fixed = fix_error_cells(rows)
    matched = sorted(mapping.values())
    missing = [column for column in ALL_COLUMNS if column not in matched]

    print(f"파일     : {path}")
    print(f"시트     : {sheet_name}")
    print(f"헤더 행  : {header_row}")
    print(f"인식 컬럼: {len(matched)}개 {matched}")
    if missing:
        print(f"미인식   : {missing} (빈 값으로 저장됨)")
    if fixed:
        print(f"수식오류 보정: {fixed}건 (계약분류를 최빈값으로 채움)")
    print(f"데이터 행: {len(rows)}건")
    for sample in rows[:5]:
        print(
            "  -",
            sample["contract_date"] or "(일자없음)",
            "|",
            sample["client"][:14],
            "|",
            sample["project_name"][:32],
            "|",
            sample["contract_amount"],
        )

    if not args.commit:
        print("\n미리보기만 했습니다. 실제 입력은 --commit 을 붙여 다시 실행하세요.")
        return

    if not rows:
        print("\n입력할 행이 없습니다.")
        return

    inserted = insert_rows(rows)
    print(f"\nbit_history_rows 에 {inserted}건 추가 완료.")


if __name__ == "__main__":
    main()
